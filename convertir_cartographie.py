#!/usr/bin/env python3
"""Convertit la feuille « Cartographie » d'un classeur Excel (.xlsm, .xlsx) en
fichier CSV de chargement en masse pour l'outil d'analyse réglementaire.

Usage :
  python3 convertir_cartographie.py CLASSEUR.xlsm [CLASSEUR2.xlsm ...]
      écrit un CSV par classeur, à côté de lui : CLASSEUR.csv ;
  python3 convertir_cartographie.py CLASSEUR.xlsm [...] -o SORTIE.csv
      réunit tous les classeurs dans un seul CSV.

Principe de la feuille source : un article s'étend sur plusieurs lignes. La
première porte la référence et le fabricant ; les suivantes (référence et
fabricant vides) ajoutent chacune une substance.

Reprise :
  Fabricant/Fournisseur   -> Nom fabricant
  Référence Fabricant     -> Réf fabricant
  Nom substance, N°CAS, N°CE, Présence ou concentration (%)
                          -> Substances (une par ligne : nom | CAS | CE | %)
Le reste (Référence ID, désignation, pays, analyses calculées par le classeur,
justificatifs, code déchet, niveaux d'ensemble) n'est pas repris. Les colonnes
de conformité et de certificats sont présentes mais vides.

Un même article (fabricant + réf, majuscules et espaces indifférents) présent
plusieurs fois est fusionné : ses substances sont réunies, sans doublon.
Aucune valeur n'est corrigée : n° CAS, n° CE, pourcentages sont recopiés tels
quels, l'analyse du fichier dans l'outil signalera ce qui est à reprendre.

Le CSV est en UTF-8 avec marque d'ordre, séparateur « ; ». Les avertissements
(lignes de substance sans article, article sans réf ou sans fabricant,
fusions) sont affichés et écrits dans SORTIE_avertissements.txt.

Demande Python 3.8+ et openpyxl (pip install openpyxl).
"""

import argparse
import csv
import sys
import warnings
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Il manque la bibliothèque openpyxl : pip install openpyxl")

FEUILLE = "Cartographie"
LIGNES_EN_TETE = 10  # les en-têtes sont cherchés dans ces premières lignes

# En-têtes de la source (comparés sans majuscules, accents gardés, espaces réduits).
SOURCE = {
    "ref": "référence fabricant",
    "fabricant": "fabricant/fournisseur",
    "nom": "nom substance (s) ou composé(s)",
    "cas": "n°cas",
    "ce": "n°ce",
    "pct": "présence ou concentration sur article (%)",
}
OBLIGATOIRES = ("ref", "fabricant", "nom")

COLONNES = [
    "Nom fabricant", "Réf fabricant", "ESN fabricant", "Substances",
    "Conformité REACH", "Certificat REACH", "Date déclaration fabricant",
    "Conformité RoHS", "Certificat RoHS", "Commentaire RoHS", "Conformité CE",
    "Directives applicables", "Certificat CE", "Validé le", "Statut actuel",
]
LIGNES_MAX_IMPORT = 20_000


def normaliser(texte):
    return " ".join(str(texte).split()).lower()


def texte(cellule):
    """Valeur d'une cellule en texte, comme Excel l'affiche pour les nombres entiers."""
    v = cellule.value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        # Format « 000000 » : Excel affiche les zéros de tête, on les garde.
        fmt = cellule.number_format or ""
        if fmt and set(fmt) == {"0"}:
            return str(v).zfill(len(fmt))
        return str(v)
    return " ".join(str(v).split())


def pourcentage(cellule):
    """Pourcentage en texte français (« 0,5 »). Une cellule au format % (0,05) donne 5."""
    v = cellule.value
    if isinstance(v, (int, float)):
        if "%" in (cellule.number_format or ""):
            v = v * 100
        return ("%.6f" % v).rstrip("0").rstrip(".").replace(".", ",")
    return texte(cellule)


def trouver_colonnes(feuille):
    """Ligne d'en-têtes et colonne (indice 0) de chaque donnée reprise."""
    lignes = feuille.iter_rows(min_row=1, max_row=LIGNES_EN_TETE, values_only=True)
    for numero, ligne in enumerate(lignes, start=1):
        trouvees = {}
        for i, v in enumerate(ligne):
            if v is None:
                continue
            entete = normaliser(v)
            for cle, libelle in SOURCE.items():
                if entete == libelle and cle not in trouvees:
                    trouvees[cle] = i
        if all(c in trouvees for c in OBLIGATOIRES):
            return numero, trouvees
    manquants = ", ".join("« %s »" % SOURCE[c] for c in OBLIGATOIRES)
    raise ValueError("en-têtes introuvables dans les %d premières lignes (il faut %s)"
                     % (LIGNES_EN_TETE, manquants))


def lire(chemin, articles, avertir):
    """Ajoute à `articles` (clé -> article) ceux de la feuille Cartographie."""
    with warnings.catch_warnings():  # dessins, zones d'impression : sans effet ici
        warnings.simplefilter("ignore")
        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    if FEUILLE not in classeur.sheetnames:
        raise ValueError("pas de feuille « %s »" % FEUILLE)
    feuille = classeur[FEUILLE]
    ligne_entetes, col = trouver_colonnes(feuille)

    def valeur(ligne, cle, conversion=texte):
        i = col.get(cle)
        return conversion(ligne[i]) if i is not None and i < len(ligne) else ""

    courant = None
    lus = 0
    lignes = feuille.iter_rows(min_row=ligne_entetes + 1)
    for numero, ligne in enumerate(lignes, start=ligne_entetes + 1):
        ref, fabricant = valeur(ligne, "ref"), valeur(ligne, "fabricant")
        if ref or fabricant:
            lus += 1
            if not ref or not fabricant:
                avertir("%s, ligne %d : article sans %s (« %s »)"
                        % (chemin.name, numero, "réf" if not ref else "fabricant",
                           ref or fabricant))
            cle = (normaliser(fabricant), normaliser(ref))
            courant = articles.get(cle)
            if courant is None:
                courant = {"fabricant": fabricant, "ref": ref, "substances": [],
                           "origine": "%s ligne %d" % (chemin.name, numero)}
                articles[cle] = courant
            else:
                avertir("%s, ligne %d : %s %s déjà vu (%s), substances réunies"
                        % (chemin.name, numero, fabricant, ref, courant["origine"]))
        substance = [valeur(ligne, "nom"), valeur(ligne, "cas"), valeur(ligne, "ce"),
                     valeur(ligne, "pct", pourcentage)]
        if not any(substance):
            continue
        if courant is None:
            avertir("%s, ligne %d : substance « %s » avant le premier article, ignorée"
                    % (chemin.name, numero, substance[0] or substance[1]))
            continue
        if substance not in courant["substances"]:
            courant["substances"].append(substance)
    classeur.close()
    return lus


def cellule_substances(substances):
    lignes = []
    for s in substances:
        while s and not s[-1]:  # champs vides de fin omis
            s = s[:-1]
        lignes.append(" | ".join(s))
    return "\n".join(lignes)


def ecrire(chemin, articles):
    with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
        sortie = csv.writer(f, delimiter=";", lineterminator="\r\n")
        sortie.writerow(COLONNES)
        for a in articles.values():
            ligne = [""] * len(COLONNES)
            ligne[0], ligne[1] = a["fabricant"], a["ref"]
            ligne[3] = cellule_substances(a["substances"])
            sortie.writerow(ligne)


def convertir(entrees, sortie):
    articles, avertissements = {}, []
    lus = 0
    for chemin in entrees:
        try:
            lus += lire(chemin, articles, avertissements.append)
        except Exception as e:  # fichier illisible, feuille ou en-têtes absents
            sys.exit("%s : %s" % (chemin, e))
    ecrire(sortie, articles)
    nb_substances = sum(len(a["substances"]) for a in articles.values())
    print("%s : %d article(s) (%d lu(s)), %d substance(s)"
          % (sortie, len(articles), lus, nb_substances))
    if len(articles) > LIGNES_MAX_IMPORT:
        avertissements.append("%d articles : l'outil en charge %d au plus par fichier, "
                              "à découper" % (len(articles), LIGNES_MAX_IMPORT))
    rapport = sortie.with_name(sortie.stem + "_avertissements.txt")
    if avertissements:
        rapport.write_text("\n".join(avertissements) + "\n", encoding="utf-8")
        print("  %d avertissement(s), détail dans %s" % (len(avertissements), rapport))
        for a in avertissements[:10]:
            print("  - " + a)
        if len(avertissements) > 10:
            print("  - …")
    elif rapport.exists():
        rapport.unlink()  # reste d'une conversion précédente


def main():
    parser = argparse.ArgumentParser(
        description="Feuille Cartographie (.xlsm) -> CSV de chargement en masse.")
    parser.add_argument("classeurs", nargs="+", type=Path, help="fichiers .xlsm ou .xlsx")
    parser.add_argument("-o", "--sortie", type=Path,
                        help="un seul CSV pour tous les classeurs")
    args = parser.parse_args()
    if args.sortie:
        convertir(args.classeurs, args.sortie)
    else:
        for chemin in args.classeurs:
            convertir([chemin], chemin.with_suffix(".csv"))


if __name__ == "__main__":
    main()
